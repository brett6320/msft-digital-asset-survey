#!/usr/bin/env python3
"""Generate templates/inventory-template.xlsx (needs openpyxl)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent.parent / "templates" / "inventory-template.xlsx"

COLUMNS = [
    ("AppID", 14), ("CampaignID", 14), ("AppCatID", 14), ("ApplicationName", 30),
    ("Criticality", 11), ("OwnerName", 22), ("OwnerEmail", 30), ("Status", 12),
    ("FormsUrl", 50), ("SentDate", 21), ("LastReminder", 21), ("ReminderCount", 14),
    ("ResponseID", 12), ("ResponseEmail", 30), ("CompletedDate", 21), ("LastError", 60),
]
STATUSES = ["Pending", "Sent", "Completed", "Escalated", "Error", "Excluded"]
# Columns the flows write timestamps or IDs into are stored as text so the
# ISO 8601 values round-trip unchanged.
TEXT_COLUMNS = {"AppID", "CampaignID", "AppCatID", "SentDate", "LastReminder",
                "ResponseID", "CompletedDate"}
MAX_ROWS = 5000

wb = Workbook()
ws = wb.active
ws.title = "Inventory"
ws.append([c for c, _ in COLUMNS])
ws.append(["EXAMPLE-001", "2026Q4", "CAT-0001", "Example Application", "L3",
           "Example Owner", "owner@example.com", "Excluded",
           '=Config!$B$2&"&"&Config!$B$3&"="&A2', "", "", 0, "", "", "", ""])
for i, (name, width) in enumerate(COLUMNS, start=1):
    letter = ws.cell(row=1, column=i).column_letter
    ws.column_dimensions[letter].width = width
    if name in TEXT_COLUMNS:
        for row in range(2, MAX_ROWS + 2):
            ws.cell(row=row, column=i).number_format = "@"
ws.freeze_panes = "A2"

table = Table(displayName="Inventory", ref=f"A1:{ws.cell(row=2, column=len(COLUMNS)).coordinate}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)

col = {name: ws.cell(row=1, column=i).column_letter for i, (name, _) in enumerate(COLUMNS, start=1)}
status_dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=False)
status_dv.add(f"{col['Status']}2:{col['Status']}{MAX_ROWS + 1}")
crit_dv = DataValidation(type="list", formula1='"L1,L2,L3,L4,L5"', allow_blank=False)
crit_dv.add(f"{col['Criticality']}2:{col['Criticality']}{MAX_ROWS + 1}")
ws.add_data_validation(status_dv)
ws.add_data_validation(crit_dv)

cfg = wb.create_sheet("Config")
cfg.append(["Setting", "Value", "Notes"])
cfg.append(["FormBaseUrl", "https://forms.office.com/Pages/ResponsePage.aspx?id=REPLACE_WITH_FORM_ID",
            "The form's response link, up to and including id=..."])
cfg.append(["AppIdQuestionId", "rREPLACE_WITH_QUESTION_ID",
            "The r... key for the AppID question, taken from a pre-filled link"])
cfg.append([])
cfg.append(["FormsUrl formula", None,
            'In Inventory, FormsUrl = Config!$B$2&"&"&Config!$B$3&"="&[AppID]. Fill it down for every row.'])
for c in ("A", "B", "C"):
    cfg.column_dimensions[c].width = {"A": 20, "B": 70, "C": 80}[c]
for cell in cfg[1]:
    cell.font = Font(bold=True)

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}")
